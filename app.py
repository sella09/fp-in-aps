import streamlit as st
import pdfplumber
import re
import pandas as pd
from io import BytesIO
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Ekstraktor Faktur Pajak", layout="wide")
st.title("📄 Ekstraktor Faktur Pajak Indonesia")
st.caption("Upload PDF faktur → Excel berisi Rekap + 1 sheet per faktur (Header + Detail Barang).")


def parse_angka(s):
    if not s:
        return None
    s = str(s).strip().replace("Rp", "").replace(" ", "")
    s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def safe_sheet_name(name, used):
    base = re.sub(r"[\\/*?:\[\]]", "_", str(name))[:28].strip()
    if not base or base[0] in "[]:*?/\\":
        base = "Faktur"
    sheet = base
    i = 1
    while sheet in used:
        suffix = f"_{i}"
        sheet = (base[:28] + suffix)[:31]
        i += 1
    used.add(sheet)
    return sheet


def ekstrak_header(teks):
    data = {}
    m = re.search(r"Nomor Seri Faktur Pajak:\s*(\d+)", teks)
    data["nomor_seri"] = m.group(1) if m else None

    npwp_all = re.findall(r"NPWP\s*:\s*(\d{15,16})", teks)
    data["npwp_penjual"] = npwp_all[0] if len(npwp_all) > 0 else None
    data["npwp_pembeli"] = npwp_all[1] if len(npwp_all) > 1 else None

    m = re.search(r"Pengusaha Kena Pajak:\s*\n?\s*Nama\s*:\s*(.+)", teks)
    data["nama_penjual"] = m.group(1).strip() if m else None

    m = re.search(r"Pembeli Barang Kena Pajak/Penerima Jasa Kena Pajak:\s*\n?\s*Nama\s*:\s*(.+)", teks)
    data["nama_pembeli"] = m.group(1).strip() if m else None

    m = re.search(r"Pengusaha Kena Pajak:.*?Alamat\s*:\s*(.+)", teks, re.DOTALL)
    data["alamat_penjual"] = m.group(1).strip().split("\n")[0] if m else None

    m = re.search(r"Pembeli Barang Kena Pajak/Penerima Jasa Kena Pajak:.*?Alamat\s*:\s*(.+)", teks, re.DOTALL)
    data["alamat_pembeli"] = m.group(1).strip().split("\n")[0] if m else None

    m = re.search(r"Dasar Pengenaan Pajak\s*([\d.,]+)", teks)
    data["dpp"] = parse_angka(m.group(1)) if m else None

    m = re.search(r"Jumlah PPN[^0-9]*([\d.,]+)", teks, re.DOTALL)
    data["ppn"] = parse_angka(m.group(1)) if m else None

    m = re.search(r"Jumlah\s+P+[nN]?[bB]?[mM]?[^0-9]*([\d.,]+)", teks)
    data["ppnbm"] = parse_angka(m.group(1)) if m else None

    m = re.search(r"Harga Jual / Penggantian / Uang Muka / Termin\s*([\d.,]+)", teks)
    data["harga_jual_total"] = parse_angka(m.group(1)) if m else None

    m = re.search(r"([A-Z][A-Z\.\s]+?),\s*(\d{1,2}\s+\w+\s+\d{4})", teks)
    data["tanggal"] = m.group(2) if m else None

    return data


def ekstrak_barang(teks):
    """Ekstrak barang dari format multi-baris."""
    items = []

    # Buang header tabel
    teks_bersih = re.sub(r"Kode\s+Harga Jual.*?\(Rp\)", "", teks, flags=re.DOTALL)
    teks_bersih = re.sub(r"No\.?\s*Kode\s*Barang.*?Jasa Kena Pajak", "", teks_bersih, flags=re.DOTALL)

    # === FORMAT A: dengan "- Part No" (SDLG style) ===
    pola = re.compile(
        r"(?P<nama>[A-Za-z][A-Za-z0-9\s\.\-/()%]*?)\s*-\s*Part No\s*:\s*(?P<part>\S+)\s*\n"
        r"\s*Rp\s*(?P<harga>[\d.,]+)\s*x\s*(?P<qty>[\d.,]+)\s*(?P<satuan>\w+)\s*\n"
        r"\s*(?P<no>\d{1,3})\s+(?P<kode>\d{6})\s+(?P<subtotal>[\d.,]+)",
        re.MULTILINE
    )
    for m in pola.finditer(teks_bersih):
        items.append({
            "No": m.group("no"),
            "Kode": m.group("kode"),
            "Nama Barang": m.group("nama").strip(),
            "Part No": m.group("part").strip(),
            "Harga Satuan": parse_angka(m.group("harga")),
            "Qty": parse_angka(m.group("qty")),
            "Satuan": m.group("satuan").strip(),
            "Subtotal": parse_angka(m.group("subtotal")),
        })

    if items:
        return items

    # === FORMAT B: tanpa "- Part No" ===
    pola_b = re.compile(
        r"(?P<nama>[A-Z][A-Za-z0-9\s\.\-/()%]*?)\s*\n"
        r"\s*Rp\s*(?P<harga>[\d.,]+)\s*x\s*(?P<qty>[\d.,]+)\s*(?P<satuan>\w+)\s*\n"
        r"\s*(?P<no>\d{1,3})\s+(?P<kode>\d{6})\s+(?P<subtotal>[\d.,]+)",
        re.MULTILINE
    )
    for m in pola_b.finditer(teks_bersih):
        nama = m.group("nama").strip()
        if any(x in nama for x in ["Harga Jual", "Nama Barang", "Kode Barang", "Kode Harga"]):
            continue
        items.append({
            "No": m.group("no"),
            "Kode": m.group("kode"),
            "Nama Barang": nama,
            "Part No": "",
            "Harga Satuan": parse_angka(m.group("harga")),
            "Qty": parse_angka(m.group("qty")),
            "Satuan": m.group("satuan").strip(),
            "Subtotal": parse_angka(m.group("subtotal")),
        })

    return items


def proses_pdf(file):
    with pdfplumber.open(file) as pdf:
        teks = "\n".join(page.extract_text() or "" for page in pdf.pages)
    header = ekstrak_header(teks)
    header["nama_file"] = file.name
    barang = ekstrak_barang(teks)
    return header, barang


def pastikan_sheet(writer, sheet):
    if sheet not in writer.sheets:
        pd.DataFrame().to_excel(writer, sheet_name=sheet)


def tulis_df(writer, df, sheet, startrow=0, startcol=0,
             kolom_angka=None, header_fill="D9E1F2"):
    pastikan_sheet(writer, sheet)
    df.to_excel(writer, sheet_name=sheet, index=False,
                startrow=startrow, startcol=startcol)
    ws = writer.sheets[sheet]
    kolom_angka = kolom_angka or []

    for col_name in kolom_angka:
        if col_name not in df.columns:
            continue
        idx = df.columns.get_loc(col_name) + startcol + 1
        letter = get_column_letter(idx)
        for row in range(startrow + 2, startrow + 2 + len(df)):
            ws[f"{letter}{row}"].number_format = "#,##0"

    for col in range(startcol + 1, startcol + 1 + len(df.columns)):
        c = ws.cell(row=startrow + 1, column=col)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=header_fill)
        c.alignment = Alignment(horizontal="center", vertical="center")

    return startrow + len(df) + 1


def tulis_judul(writer, sheet, text, row):
    pastikan_sheet(writer, sheet)
    ws = writer.sheets[sheet]
    c = ws.cell(row=row + 1, column=1, value=text)
    c.font = Font(bold=True, size=12)
    return row + 2


uploaded_files = st.file_uploader(
    "Upload file PDF faktur pajak (bisa banyak)",
    type=["pdf"], accept_multiple_files=True,
)

if uploaded_files:
    if st.button("🚀 Proses", type="primary"):
        hasil = []
        progress = st.progress(0)
        status = st.empty()

        for i, f in enumerate(uploaded_files):
            status.text(f"Memproses: {f.name} ({i+1}/{len(uploaded_files)})")
            try:
                h, b = proses_pdf(f)
                hasil.append({"header": h, "barang": b})
            except Exception as e:
                st.warning(f"Gagal memproses {f.name}: {e}")
            progress.progress((i + 1) / len(uploaded_files))
        status.text("✅ Selesai!")

        total = len(hasil)
        total_barang = sum(len(r["barang"]) for r in hasil)

        c1, c2 = st.columns(2)
        c1.metric("Total Faktur", total)
        c2.metric("Total Barang", total_barang)

        for r in hasil:
            h = r["header"]
            st.markdown("---")
            st.subheader(f"📄 {h.get('nomor_seri') or h.get('nama_file')}")
            st.caption(f"Penjual: {h.get('nama_penjual')} | Pembeli: {h.get('nama_pembeli')} | Tanggal: {h.get('tanggal')}")
            st.markdown("**Detail Barang**")
            df_b = pd.DataFrame(r["barang"])
            if not df_b.empty:
                st.dataframe(df_b, use_container_width=True)
            else:
                st.warning("⚠️ Tidak ada barang yang terekstrak dari PDF ini.")

        # ================= EXPORT EXCEL =================
        buffer = BytesIO()
        used_names = set()

        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            # ---------- SHEET REKAP ----------
            rekap_sheet = safe_sheet_name("Rekap", used_names)
            pd.DataFrame([["REKAP FAKTUR PAJAK"]]).to_excel(
                writer, sheet_name=rekap_sheet, index=False, header=False
            )
            writer.sheets[rekap_sheet].cell(row=1, column=1).font = Font(bold=True, size=14)

            rekap_rows = []
            for r in hasil:
                h = r["header"]
                total_item = sum(b.get("Subtotal") or 0 for b in r["barang"])
                rekap_rows.append({
                    "Nomor Seri": h.get("nomor_seri"),
                    "Tanggal": h.get("tanggal"),
                    "Nama Penjual": h.get("nama_penjual"),
                    "NPWP Penjual": h.get("npwp_penjual"),
                    "Nama Pembeli": h.get("nama_pembeli"),
                    "NPWP Pembeli": h.get("npwp_pembeli"),
                    "Jumlah Item": len(r["barang"]),
                    "Total Item": total_item,
                    "DPP": h.get("dpp"),
                    "PPN": h.get("ppn"),
                    "PPnBM": h.get("ppnbm"),
                    "Harga Jual Total": h.get("harga_jual_total"),
                    "Nama File": h.get("nama_file"),
                })
            df_rekap = pd.DataFrame(rekap_rows)

            total_row = pd.DataFrame([{
                "Nomor Seri": "TOTAL", "Tanggal": "", "Nama Penjual": "",
                "NPWP Penjual": "", "Nama Pembeli": "", "NPWP Pembeli": "",
                "Jumlah Item": df_rekap["Jumlah Item"].sum() if not df_rekap.empty else 0,
                "Total Item": df_rekap["Total Item"].sum() if not df_rekap.empty else 0,
                "DPP": df_rekap["DPP"].sum() if not df_rekap.empty else 0,
                "PPN": df_rekap["PPN"].sum() if not df_rekap.empty else 0,
                "PPnBM": df_rekap["PPnBM"].sum() if not df_rekap.empty else 0,
                "Harga Jual Total": df_rekap["Harga Jual Total"].sum() if not df_rekap.empty else 0,
                "Nama File": "",
            }])
            df_rekap_full = pd.concat([df_rekap, total_row], ignore_index=True)

            tulis_df(
                writer, df_rekap_full, rekap_sheet, startrow=2,
                kolom_angka=["Total Item", "DPP", "PPN", "PPnBM",
                             "Harga Jual Total", "Jumlah Item"],
            )
            ws = writer.sheets[rekap_sheet]
            for col in range(1, len(df_rekap_full.columns) + 1):
                ws.cell(row=2 + len(df_rekap_full), column=col).font = Font(bold=True)

            # ---------- 1 SHEET PER FAKTUR ----------
            for r in hasil:
                h = r["header"]
                raw_name = h.get("nomor_seri") or h.get("nama_file", "Faktur")
                sheet = safe_sheet_name(str(raw_name), used_names)

                df_info = pd.DataFrame([
                    ["Nomor Seri", h.get("nomor_seri")],
                    ["Tanggal", h.get("tanggal")],
                    ["NPWP Penjual", h.get("npwp_penjual")],
                    ["Nama Penjual", h.get("nama_penjual")],
                    ["Alamat Penjual", h.get("alamat_penjual")],
                    ["NPWP Pembeli", h.get("npwp_pembeli")],
                    ["Nama Pembeli", h.get("nama_pembeli")],
                    ["Alamat Pembeli", h.get("alamat_pembeli")],
                    ["DPP", h.get("dpp")],
                    ["PPN", h.get("ppn")],
                    ["PPnBM", h.get("ppnbm")],
                    ["Harga Jual Total", h.get("harga_jual_total")],
                ], columns=["Field", "Nilai"])

                row = 0
                row = tulis_judul(writer, sheet, "HEADER FAKTUR", row)
                row = tulis_df(writer, df_info, sheet, startrow=row,
                               kolom_angka=["Nilai"])

                df_barang = pd.DataFrame(r["barang"])
                row += 1
                row = tulis_judul(writer, sheet, "DETAIL BARANG", row)
                row = tulis_df(writer, df_barang, sheet, startrow=row,
                               kolom_angka=["Harga Satuan", "Qty", "Subtotal"])

                ws = writer.sheets[sheet]
                for col_cells in ws.columns:
                    max_len = 0
                    letter = col_cells[0].column_letter
                    for cell in col_cells:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    ws.column_dimensions[letter].width = min(max_len + 2, 60)

        st.download_button(
            label="⬇️ Download Excel (Rekap + 1 sheet per faktur)",
            data=buffer.getvalue(),
            file_name="faktur_pajak.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

else:
    st.info("Silakan upload file PDF faktur pajak terlebih dahulu.")
